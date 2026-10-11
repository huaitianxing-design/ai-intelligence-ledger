#!/usr/bin/env python3
# AI情报勾稽看板·增量爬取管线
# 直连源：虎嗅(m站)+量子位(官网)+Deeptech(公开API)；公众号源：搜狗微信索引(12源) -> 还原微信直链 -> 去重 -> 自动分级 -> 增量写入看板与Excel底稿
import io, json, os, re, subprocess, sys, time, urllib.request, html as H
from datetime import date

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FILES = [ROOT + '/index.html', ROOT + '/offline/AI情报勾稽看板_离线版.html']
XLSX = ROOT + '/data/AI情报勾稽底稿.xlsx'
UA = 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/126.0 Safari/537.36'
UA_IPHONE = 'Mozilla/5.0 (iPhone; CPU iPhone OS 17_0 like Mac OS X) AppleWebKit/605.1.15'
CK = os.getenv('LEDGER_SOGOU_COOKIE_JAR', os.path.join(ROOT, '.sogou-cookies.txt'))
AI_KW = ['AI', '英伟达', 'OpenAI', '大模型', '智能体', 'Agent', '算力', '芯片', 'Gemini', 'GPT', 'Claude',
         '智谱', 'MiniMax', '豆包', '谷歌', 'Meta', '微软', '亚马逊', '财报', '资本开支', '通义', '混元', '文心']
B_KW = ['英伟达', '谷歌', 'Google', 'Meta', '微软', '亚马逊', '阿里', '腾讯', '快手', '美团', '智谱', 'MiniMax',
        '资本开支', 'Capex', '财报', '营收', 'ARR', '广告收入', 'OpenAI', 'Anthropic']
# 12个公众号源（搜狗微信通道）
WX_SOURCES = ['Deeptech深科技', '量子位', '海豚研究', '虎嗅', '腾讯科技',
              '华尔街博客', 'bedrock', '404K', '2030FY', '我有充足的时间', '地平线全球策略', '星辰大海的边界']
MAX_NEW = 50

def fetch(url, timeout=15, ua=UA, referer=None):
    h = {'User-Agent': ua, 'Accept-Language': 'zh-CN,zh;q=0.9'}
    if referer: h['Referer'] = referer
    req = urllib.request.Request(url, headers=h)
    return urllib.request.urlopen(req, timeout=timeout).read().decode('utf-8', 'ignore')

def clean(t):
    t = re.sub(r'<(script|style).*?</\1>', '', t, flags=re.S)
    t = re.sub(r'<[^>]+>', '', t)
    return re.sub(r'\s+', '', H.unescape(t))

def num_anchors(body, n=3):
    out = []
    for m in re.finditer(r'[^，。；、!?]{0,30}\d[\d,.，%亿万倍亿美元条家款个KkMm]{0,18}[^，。；、!?]{0,20}', body):
        s = m.group(0).strip()
        if 6 < len(s) < 60 and s not in out:
            out.append(s)
        if len(out) >= n: break
    return out or ['要点待人工补充']

def grade(title):
    return ('B', 4) if any(k in title for k in B_KW) else ('C', 3)

def topic_of(title):
    return 'T-CAPEX-001' if any(k in title for k in ['资本开支', 'Capex', '算力', '营收', '财报', 'ARR']) else 'T-IND-001'

# ---------- 直连源 ----------
def huxiu_items():
    try: s = fetch('https://m.huxiu.com/article/')
    except Exception as e: print('huxiu feed fail', e); return []
    s = s.replace('<!---->', '')
    out = []
    for m in re.finditer(r'href="(https://m\.huxiu\.com/article/(\d+)\.html)[^"]*"[^>]*class="item"(.*?)</a>', s, re.S):
        url, aid, seg = m.group(1), m.group(2), m.group(3)
        tm = re.search(r'title-wrap[^>]*>\s*([^<]+?)\s*</h2>', seg)
        if tm: out.append({'url': url, 'title': H.unescape(tm.group(1)).strip(), 'src': '虎嗅', 'cls': 'AI垂直类'})
    return out

def deeptech_items():
    """Deeptech深科技官网直连：apii.web.mittrchina.com 公开API，翻3页"""
    hd = {'User-Agent': UA, 'Origin': 'https://www.mittrchina.com', 'Referer': 'https://www.mittrchina.com/'}
    items = []
    for pg in (1, 2, 3):
        try:
            req = urllib.request.Request('https://apii.web.mittrchina.com/information/index?page=%d&limit=20' % pg, headers=hd)
            d = json.loads(urllib.request.urlopen(req, timeout=12).read().decode('utf-8', 'ignore'))
            items += (d.get('data') or {}).get('items') or []
        except Exception as e:
            print('deeptech feed fail p%d' % pg, e)
    if not items: return []
    out, seen = [], set()
    for x in items:
        if x.get('id') in seen: continue
        seen.add(x.get('id'))
        body, ct = '', ''
        try:
            req2 = urllib.request.Request('https://apii.web.mittrchina.com/information/details?id=%s' % x.get('id'), headers=hd)
            det = (json.loads(urllib.request.urlopen(req2, timeout=12).read().decode('utf-8', 'ignore')).get('data') or {})
            body = clean(det.get('summary') or '')[:600] if isinstance(det.get('summary'), str) else ''
            ct = str(det.get('create_time') or '')[:10]
        except Exception as e: print('deeptech detail fail', x.get('id'), e)
        out.append({'url': 'https://www.mittrchina.com/news/detail/%s' % x.get('id'),
                    'title': (x.get('name') or '').strip(), 'src': 'Deeptech深科技', 'cls': 'AI垂直类',
                    'desc': body, 'date': ct})
    return out

def qbitai_items():
    urls, seen = [], set()
    for pg in ('', 'page/2/', 'page/3/'):
        try:
            s = fetch('https://www.qbitai.com/' + pg)
        except Exception as e:
            print('qbitai feed fail', pg, e); continue
        for u in re.findall(r'href="(https://www\.qbitai\.com/20\d\d/\d{2}/\d+\.html)"', s):
            if u not in seen: seen.add(u); urls.append(u)
    out = []
    for u in urls[:20]:
        try:
            p = fetch(u)
            t = re.search(r'<title>([^<]+)</title>', p)
            if not t: continue
            out.append({'url': u, 'title': H.unescape(t.group(1)).split('-量子位')[0].strip(), 'src': '量子位',
                        'cls': 'AI垂直类', 'desc': clean(p)[:600]})
        except Exception as e: print('qbitai art fail', u, e)
    return out

# ---------- 搜狗微信通道（12源） ----------
def curl(url, referer=None, jar_in=True, jar_out=False, ua=UA):
    cmd = ['curl', '-s', '-m', '15', '-A', ua]
    if jar_in and os.path.exists(CK): cmd += ['-b', CK]
    if jar_out: cmd += ['-c', CK]
    if referer: cmd += ['-e', referer]
    cmd.append(url)
    return subprocess.run(cmd, capture_output=True, text=True).stdout

def sogou_items():
    out, got_urls = [], set()
    for src in WX_SOURCES:
        queries = [src + ' AI', src + ' 财报 算力']
        got = 0
        for qi, qkw in enumerate(queries):
            if got >= 4: break
            q = urllib.parse.quote(qkw)
            page = curl('https://weixin.sogou.com/weixin?type=2&query=' + q + ('&page=2' if qi == 0 and got == 0 else ''), jar_out=(not out))
            if 'antispider' in page or '验证码' in page:
                print('sogou antispider at', src); return out
            items = re.findall(r'<a[^>]+href="(/link\?url=[^"]+)"[^>]*uigs="article_title_(\d+)"[^>]*>(.*?)</a>', page, re.S)
            for h, i, t in items:
                if got >= 4: break
                title = H.unescape(re.sub(r'<[^>]+>', '', t)).replace('&nbsp;', ' ').strip()
                acc_m = re.search(r'uigs="article_title_' + i + r'".{0,900}?class="all-time-y2">([^<]+)</span>', page, re.S)
                acc = H.unescape(acc_m.group(1)).strip() if acc_m else ''
                if src not in acc and acc not in src:  # 账号归属不符，跳过防张冠李戴
                    continue
                u = 'https://weixin.sogou.com' + urllib.parse.quote(H.unescape(h), safe='/?:=&')
                if u in got_urls: continue
                link_page = curl(u, referer='https://weixin.sogou.com/weixin?type=2&query=' + q)
                real = ''.join(re.findall(r"url\s*\+=\s*'([^']*)'", link_page))
                if not real: continue
                if real.split('&timestamp=')[0] in got_urls: continue
                got_urls.add(u); got_urls.add(real.split('&timestamp=')[0])
                wx = curl(real, ua=UA_IPHONE)
                if '该内容已被发布者删除' in wx or '此内容因违规无法查看' in wx or 'js_content' not in wx:
                    continue  # 被删/壳页，不入库
                mt = re.search(r'var msg_title = "([^"]+)"|var msg_title = \'([^\']+)\'' + r'|<h1[^>]*class="rich_media_title"[^>]*>\s*([^<]+)', wx)
                title2 = H.unescape(next((x for x in (mt.groups() if mt else ()) if x), title)).strip()
                body = clean(wx)[:600]
                out.append({'url': real, 'title': title2[:60], 'src': src, 'cls': 'AI垂直类' if src in WX_SOURCES[:5] else '财经分析',
                            'desc': body, 'wx': True})
                got += 1
                time.sleep(2)
            time.sleep(3)
        time.sleep(3)
    return out

import urllib.parse

def main():
    today = date.today().isoformat()
    s = io.open(FILES[0], encoding='utf-8').read()
    ci = s.index('const cards=['); cj = s.index('];', ci)
    cards = json.loads(s[ci + len('const cards='):cj].rstrip().rstrip(',') + ']')
    have_urls = {x.get('url', '').split('&timestamp=')[0] for x in cards}
    have_titles = {re.sub(r'\W', '', x['title']) for x in cards}
    nxt = max([int(x['id'].split('-')[-1]) for x in cards if x['id'].startswith('INT-')] or [0]) + 1

    cands = huxiu_items() + qbitai_items() + deeptech_items() + sogou_items()
    print('candidates:', len(cands))
    new = []
    for c in cands:
        if c['url'].split('&timestamp=')[0] in have_urls: continue
        if re.sub(r'\W', '', c['title']) in have_titles: continue
        if not any(k in c['title'] for k in AI_KW): continue
        body = c.get('desc') or ''
        if c['src'] == '虎嗅':
            try:
                p = fetch(c['url'])
                d = re.search(r'"publish_time[^"]*"\s*:\s*"([^"]+)"', p) or re.search(r'(\d{4}-\d{2}-\d{2})', p)
                c['date'] = (d.group(1)[:10] if d else today)
                body = clean(p)[:600]
            except Exception: c['date'] = today
        else:
            c['date'] = c.get('date') or today
        g, sc = grade(c['title'])
        note = '自动爬取·待人工复核'
        if c.get('wx'): note += '；微信链接受签名时效限制，失效后按标题在搜狗微信重查'
        new.append({'id': 'INT-%s-%03d' % (today.replace('-', ''), nxt), 'date': c['date'], 'source': c['src'],
                    'title': c['title'][:60], 'url': c['url'], 'grade': g, 'topic': topic_of(c['title']), 'company': '行业',
                    'score': sc, 'anchors': num_anchors(body), 'summary': (body[:110] + '…') if body else '待人工补充摘要。',
                    'drill': '自动抓取条目；下钻原文并人工核对归属公司与口径。', 'status': '全文已缓存' if body else '待核验',
                    'note': note})
        nxt += 1
        have_urls.add(c['url'].split('&timestamp=')[0])
        if len(new) >= MAX_NEW: break
    print('new cards:', len(new), [x['source'] + '|' + x['title'][:16] for x in new])
    if not new:
        print('NO_NEW'); return

    block = ''.join(',\n' + json.dumps(x, ensure_ascii=False) for x in new)
    for p in FILES:
        t = io.open(p, encoding='utf-8').read()
        if new[0]['id'] in t: continue
        ci = t.index('const cards=['); cj = t.index('];', ci)
        seg = t[ci:cj]
        m = re.search(r'\}(\s*,?\s*)\Z', seg)
        ins = ci + m.start() + 1
        t = t[:ins] + block + t[ins:]
        io.open(p, 'w', encoding='utf-8').write(t)
        print('patched', p)

    import openpyxl
    wb = openpyxl.load_workbook(XLSX)
    ws = wb['03_公众号信息']
    hdr = [c.value for c in ws[1]]
    ws1 = wb['01_信息源']
    h1 = [c.value for c in ws1[1]]
    sid_map = {}
    for row in ws1.iter_rows(min_row=2, values_only=True):
        if row and len(row) > 3 and row[3]: sid_map[str(row[3]).strip()] = row[0]
    r = ws.max_row
    for x in new:
        r += 1
        row = {'情报ID': x['id'], '发布日期': x['date'], '抓取日期': today, '来源ID': sid_map.get(x['source'], ''),
               '来源名称': x['source'], '标题': x['title'], '原文链接': x['url'], '分类': x.get('cls', 'AI垂直类'),
               '主题编号': x['topic'], '公司': x['company'], '相关性评分(1-5)': x['score']}
        for k, v in enumerate(x['anchors'][:5], 1): row['数字锚点%d' % k] = v
        row.update({'全文总结': x['summary'], '下钻建议': x['drill'], '处理状态': x['status'], '备注': x['note']})
        for kk, v in row.items():
            if kk in hdr: ws.cell(row=r, column=hdr.index(kk) + 1, value=v)
    wb.save(XLSX)
    print('excel rows ->', ws.max_row)

    print('done. 新增条目已写入看板与底稿，复查后可自行 git commit。')

if __name__ == '__main__':
    main()
